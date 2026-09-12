"""表格要挂在它所在的那一节下面，不是一律甩到文档根上。

实测一份凭据技术标准：`4.1`–`4.5` 与 `5` 六个章节正文为空，它们的内容是
文末孤零零的 Table 5–10——那六张表**就是**口令规则和 MFA 强制要求，
一份凭据标准最实质的部分。原因是解析器先走完 `document.paragraphs`（不含表格），
再单独遍历 `document.tables` 一律挂到根上，表格在第几节里完全看不出来。

GRC 文档里表格常常就是要求本身（口令矩阵、SLA 表、角色矩阵），
所以表格成为**所在章节的子条款**而不是并进正文——压平会丢掉哪一列是要求、
哪一列是分类。
"""

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.parsing.docx_parser import DocxParser


def _walk(nodes, depth=0):
    for node in nodes:
        yield depth, node
        yield from _walk(node.children, depth + 1)


@pytest.fixture
def document_with_a_table_inside_a_section(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_heading("Credential Requirements", level=1)
    doc.add_paragraph("The requirements below apply to all accounts.")
    doc.add_heading("End-User Passwords", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "Length"
    table.cell(1, 1).text = "At least twelve characters"
    doc.add_heading("Exceptions", level=1)
    doc.add_paragraph("Exceptions require approval.")
    path = tmp_path / "with-table.docx"
    doc.save(path)
    return path


def test_a_table_becomes_a_child_of_the_section_it_sits_in(
    document_with_a_table_inside_a_section: Path,
):
    parsed = DocxParser().parse(document_with_a_table_inside_a_section)

    tables = [node for _, node in _walk(parsed.clauses) if node.kind == "table"]
    assert len(tables) == 1
    parents = [
        node
        for _, node in _walk(parsed.clauses)
        if any(child is tables[0] for child in node.children)
    ]
    assert [parent.heading for parent in parents] == ["End-User Passwords"]


def test_the_table_is_no_longer_a_root(document_with_a_table_inside_a_section: Path):
    parsed = DocxParser().parse(document_with_a_table_inside_a_section)

    assert [node.kind for node in parsed.clauses] == ["section", "section"]


def test_the_section_holding_a_table_is_not_reported_as_empty(
    document_with_a_table_inside_a_section: Path,
):
    """空章节会被当成漏抽——而它的内容其实就在那张表里。"""
    parsed = DocxParser().parse(document_with_a_table_inside_a_section)

    section = next(
        node for _, node in _walk(parsed.clauses) if node.heading == "End-User Passwords"
    )
    assert section.children, "这一节的内容在表里，不能表现得像一节空章节"


def test_a_table_before_any_heading_stays_at_the_root(tmp_path: Path):
    """封面上的版本历史表、缩略语表——它们不属于任何章节。"""
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Version"
    table.cell(0, 1).text = "1.0"
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Body.")
    path = tmp_path / "cover-table.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)

    assert [node.kind for node in parsed.clauses] == ["table", "section"]


def test_paragraphs_after_a_table_still_belong_to_their_section(tmp_path: Path):
    """按文档顺序遍历之后，表格不能把后面的正文吃掉或错位。"""
    doc = DocxDocument()
    doc.add_heading("Scope", level=1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "In scope"
    doc.add_paragraph("This paragraph follows the table.")
    path = tmp_path / "after-table.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)

    scope = parsed.clauses[0]
    assert "This paragraph follows the table." in scope.text
