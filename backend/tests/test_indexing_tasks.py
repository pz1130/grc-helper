from unittest.mock import AsyncMock, patch

import pytest
from sqlalchemy import select

from app.clauses.models import Clause, ClauseChunk
from app.indexing.tasks import index_document
from app.ingest.models import DocType, Document
from app.llm.routing import RoutingError


async def _document(db_session) -> Document:
    doc = Document(
        title="P", doc_type=DocType.PROCEDURE, file_hash="a" * 64,
        file_path="/x.pdf", original_filename="x.pdf",
    )
    db_session.add(doc)
    await db_session.flush()
    db_session.add(
        Clause(
            document_id=doc.id, number="1", heading="Scope", heading_path="Doc › Scope",
            citation_label="1", text="Applies to all systems.", order_index=0, level=1,
        )
    )
    await db_session.flush()
    await db_session.commit()
    return doc


@pytest.mark.asyncio
async def test_missing_embedding_provider_is_not_a_failure(db_session):
    """没配 provider 是预期状态，不是故障。

    当成失败会触发三次重试、刷屏，还会掩盖真正的上游故障。
    """
    doc = await _document(db_session)

    with patch("app.indexing.tasks.session_factory") as factory:
        factory.return_value.__aenter__.return_value = db_session
        with patch(
            "app.indexing.tasks.embed_pending",
            new=AsyncMock(side_effect=RoutingError("任务 embedding 尚未配置 provider")),
        ):
            result = await index_document({}, doc.id)

    assert result["chunks"] >= 1
    assert result["embedded"] == 0
    assert "skipped" in result


@pytest.mark.asyncio
async def test_chunks_survive_when_embedding_is_unavailable(db_session):
    """切块必须先落库：否则没配 provider 时连关键词检索都没得用。"""
    doc = await _document(db_session)

    with patch("app.indexing.tasks.session_factory") as factory:
        factory.return_value.__aenter__.return_value = db_session
        with patch(
            "app.indexing.tasks.embed_pending",
            new=AsyncMock(side_effect=RoutingError("未配置")),
        ):
            await index_document({}, doc.id)

    assert await db_session.scalar(select(ClauseChunk)) is not None


@pytest.mark.asyncio
async def test_real_upstream_failure_still_propagates(db_session):
    """上游 500 是真故障，必须抛出去让 ARQ 重试——不能和"没配"混为一谈。"""
    from app.llm.providers.base import ProviderError

    doc = await _document(db_session)

    with patch("app.indexing.tasks.session_factory") as factory:
        factory.return_value.__aenter__.return_value = db_session
        with patch(
            "app.indexing.tasks.embed_pending",
            new=AsyncMock(side_effect=ProviderError("上游返回 500", retryable=True)),
        ):
            with pytest.raises(ProviderError):
                await index_document({}, doc.id)


@pytest.mark.asyncio
async def test_parse_document_persists_document_fields_not_just_clauses(tmp_path):
    """回归：try/finally 曾写在 async with 外面。

    session 一 close，document 就成了 detached 对象，对它的赋值全被静默丢弃；
    而 persist() 新 add 的 Clause 照常落库——于是看起来"解析成功"，实际状态
    永远停在 uploaded、file_path 还指向已被删掉的 /tmp 路径，原文再也找不回来。
    单元测试里直接调 run_parse 是看不出来的，必须走 parse_document 这条真实路径。
    """
    import os

    from docx import Document as DocxDocument
    from sqlalchemy import select
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.ingest.models import DocStatus, DocType, Document
    from app.ingest.tasks import parse_document

    url = os.environ.get("TEST_DATABASE_URL", "postgresql+asyncpg://grc:grc@db:5432/grc_test")
    engine = create_async_engine(url)
    factory = async_sessionmaker(engine, expire_on_commit=False)

    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Purpose.")
    doc.add_heading("Roles and Responsibilities", level=1)
    doc.add_paragraph("IT Division.")
    path = tmp_path / "g.docx"
    doc.save(path)

    async with factory() as setup:
        document = Document(
            title="registry probe", doc_type=DocType.GUIDELINE,
            file_hash="9" * 64, file_path=str(path), original_filename="g.docx",
        )
        setup.add(document)
        await setup.commit()
        document_id = document.id

    try:
        with patch("app.ingest.tasks.session_factory", factory):
            with patch("app.ingest.tasks.enqueue" if False else "app.worker.enqueue",
                       new=AsyncMock(return_value="job")):
                await parse_document({}, document_id)

        # 关键：换一个全新 session 去读，确认字段真的落库了
        async with factory() as check:
            fresh = await check.get(Document, document_id)
            assert fresh.status is DocStatus.ACTIVE, "状态没落库——document 是 detached 的"
            assert await check.scalar(
                select(Document).where(Document.id == document_id)
            ) is not None
    finally:
        async with factory() as cleanup:
            stale = await cleanup.get(Document, document_id)
            if stale is not None:
                await cleanup.delete(stale)
                await cleanup.commit()
        await engine.dispose()
