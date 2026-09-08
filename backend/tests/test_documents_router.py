import io
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document as DocxDocument

from app.iam.models import User
from app.iam.permissions import Role
from app.iam.security import hash_password


async def _seed(db_session, role: Role, email: str) -> User:
    user = User(email=email, name=email, role=role, password_hash=hash_password("pw123456"))
    db_session.add(user)
    await db_session.flush()
    return user


async def _auth(client, email: str) -> dict[str, str]:
    response = await client.post(
        "/api/auth/login", json={"email": email, "password": "pw123456"}
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Purpose.")
    doc.add_heading("Roles and Responsibilities", level=1)
    doc.add_paragraph("IT Division.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_upload_requires_document_write(client, db_session):
    await _seed(db_session, Role.VIEWER, "v@example.com")
    headers = await _auth(client, "v@example.com")

    response = await client.post(
        "/api/documents",
        files={"file": ("g.docx", _docx_bytes())},
        data={"title": "Guideline", "doc_type": "guideline"},
        headers=headers,
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_upload_enqueues_a_parse_job(client, db_session):
    await _seed(db_session, Role.CONTRIBUTOR, "c@example.com")
    headers = await _auth(client, "c@example.com")

    with patch("app.ingest.router.enqueue", new=AsyncMock(return_value="job-1")) as job:
        response = await client.post(
            "/api/documents",
            files={"file": ("g.docx", _docx_bytes())},
            data={"title": "Guideline", "doc_type": "guideline"},
            headers=headers,
        )
    assert response.status_code == 201
    assert response.json()["status"] == "uploaded"
    job.assert_awaited_once()


@pytest.mark.asyncio
async def test_duplicate_upload_is_rejected_with_conflict(client, db_session):
    await _seed(db_session, Role.CONTRIBUTOR, "c@example.com")
    headers = await _auth(client, "c@example.com")
    payload = _docx_bytes()

    with patch("app.ingest.router.enqueue", new=AsyncMock(return_value="job-1")):
        first = await client.post(
            "/api/documents",
            files={"file": ("g.docx", payload)},
            data={"title": "G", "doc_type": "guideline"},
            headers=headers,
        )
        second = await client.post(
            "/api/documents",
            files={"file": ("g-again.docx", payload)},
            data={"title": "G", "doc_type": "guideline"},
            headers=headers,
        )
    assert first.status_code == 201
    assert second.status_code == 409
    assert "已存在" in second.json()["message"]


@pytest.mark.asyncio
async def test_unsupported_extension_is_rejected(client, db_session):
    await _seed(db_session, Role.CONTRIBUTOR, "c@example.com")
    headers = await _auth(client, "c@example.com")

    response = await client.post(
        "/api/documents",
        files={"file": ("matrix.xlsx", b"x")},
        data={"title": "M", "doc_type": "standard"},
        headers=headers,
    )
    assert response.status_code == 400
    assert "xlsx" in response.json()["message"]


@pytest.mark.asyncio
async def test_upload_is_audited(client, db_session):
    await _seed(db_session, Role.ADMIN, "a@example.com")
    headers = await _auth(client, "a@example.com")

    with patch("app.ingest.router.enqueue", new=AsyncMock(return_value="job-1")):
        await client.post(
            "/api/documents",
            files={"file": ("g.docx", _docx_bytes())},
            data={"title": "Guideline", "doc_type": "guideline"},
            headers=headers,
        )
    log = await client.get("/api/audit-log?entity_type=Document", headers=headers)
    assert any(entry["action"] == "document.upload" for entry in log.json())


@pytest.mark.asyncio
async def test_clause_tree_is_nested(client, db_session):
    from app.clauses.models import Clause
    from app.ingest.models import DocType, Document

    await _seed(db_session, Role.VIEWER, "v@example.com")
    headers = await _auth(client, "v@example.com")

    doc = Document(
        title="P",
        doc_type=DocType.PROCEDURE,
        file_hash="a" * 64,
        file_path="/x.pdf",
        original_filename="x.pdf",
    )
    db_session.add(doc)
    await db_session.flush()
    parent = Clause(
        document_id=doc.id,
        number="4",
        heading="Process",
        heading_path="Process",
        citation_label="4",
        text="",
        order_index=0,
        level=1,
    )
    db_session.add(parent)
    await db_session.flush()
    db_session.add(
        Clause(
            document_id=doc.id,
            parent_id=parent.id,
            number="4.1",
            heading="Normal",
            heading_path="Process › Normal",
            citation_label="4.1",
            text="body",
            order_index=1,
            level=2,
        )
    )
    await db_session.flush()

    response = await client.get(f"/api/documents/{doc.id}/clauses", headers=headers)
    tree = response.json()
    assert len(tree) == 1
    assert tree[0]["citation_label"] == "4"
    assert tree[0]["children"][0]["citation_label"] == "4.1"


@pytest.mark.asyncio
async def test_plain_text_fallback_recovers_a_failed_document(client, db_session):
    from app.ingest.models import DocStatus, DocType, Document

    await _seed(db_session, Role.CONTRIBUTOR, "c@example.com")
    headers = await _auth(client, "c@example.com")

    doc = Document(
        title="P",
        doc_type=DocType.PROCEDURE,
        file_hash="b" * 64,
        file_path="/x.pdf",
        original_filename="x.pdf",
        status=DocStatus.PARSE_FAILED,
        parse_error="加密的 PDF",
    )
    db_session.add(doc)
    await db_session.flush()

    response = await client.post(
        f"/api/documents/{doc.id}/plain-text",
        json={"text": "1 Scope\nThis procedure applies to all systems.\n2 Purpose\nDefine it."},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["status"] == "active"

    tree = await client.get(f"/api/documents/{doc.id}/clauses", headers=headers)
    assert [node["citation_label"] for node in tree.json()] == ["1", "2"]


@pytest.mark.asyncio
async def test_document_list_is_readable_by_viewer(client, db_session):
    await _seed(db_session, Role.VIEWER, "v@example.com")
    headers = await _auth(client, "v@example.com")
    assert (await client.get("/api/documents", headers=headers)).status_code == 200
