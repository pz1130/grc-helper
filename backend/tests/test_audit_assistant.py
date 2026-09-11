from io import BytesIO

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from sqlalchemy import select

from app.audit.citations import AnswerCitationValidator
from app.audit.history import similar_history, similarity
from app.audit.models import (
    AnswerDraft,
    AuditEngagement,
    AuditType,
    EngagementStatus,
    QuestionStatus,
)
from app.audit.service import finalize_answer, generate_answer, import_questions
from app.clauses.models import Clause
from app.controls.models import Control
from app.errors import Conflict
from app.iam.models import AuditLog, User
from app.iam.permissions import Role
from app.iam.security import create_access_token, hash_password
from app.ingest.models import DocStatus, DocType, Document
from app.llm.runner import ValidatedResult
from app.review.models import Proposal, ProposalKind, ProposalStatus
from app.review.service import Decision, decide
from app.search.schemas import SearchHit, SearchResponse


async def _user(db_session, role=Role.GRC_LEAD):
    user = User(
        email=f"{role.value}@example.com", name="User", role=role, password_hash=hash_password("pw")
    )
    db_session.add(user)
    await db_session.flush()
    return user


def _headers(user):
    return {"Authorization": f"Bearer {create_access_token(user.id, user.role)}"}


async def _engagement(db_session, user):
    row = AuditEngagement(
        name="ISO audit",
        audit_type=AuditType.EXTERNAL,
        status=EngagementStatus.PREPARING,
        created_by=user.id,
    )
    db_session.add(row)
    await db_session.flush()
    return row


@pytest.mark.asyncio
async def test_import_questions_preserves_order_and_audits(db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)

    rows = await import_questions(
        db_session, engagement.id, "First question?\n- Second question?", "en", actor=actor
    )

    assert [(row.seq, row.question_text) for row in rows] == [
        (1, "First question?"),
        (2, "Second question?"),
    ]
    audit = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_questions.import")
    )
    assert audit.after["count"] == 2


@pytest.mark.asyncio
async def test_accepting_answer_proposal_materializes_a_draft(db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (
        await import_questions(
            db_session, engagement.id, "How is access reviewed?", "en", actor=actor
        )
    )[0]
    document = Document(
        title="Policy",
        doc_type=DocType.POLICY,
        file_hash="a" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number="1",
        heading="Review",
        heading_path="Review",
        citation_label="1",
        text="Access is reviewed quarterly.",
        order_index=0,
        level=1,
    )
    control = Control(
        code="C-AUDIT", title="Access review", statement="Access is reviewed quarterly."
    )
    db_session.add_all([clause, control])
    await db_session.flush()
    payload = {
        "question_id": question.id,
        "body": "Access is reviewed quarterly.",
        "language": "en",
        "citations": [{"clause_id": clause.id, "quote": "Access is reviewed quarterly."}],
        "cited_control_ids": [control.id],
        "suggested_evidence_ids": [],
        "gap_notes": "No evidence is registered.",
        "confidence": 0.9,
    }
    proposal = Proposal(
        kind=ProposalKind.ANSWER,
        payload=payload,
        citations=payload["citations"],
        confidence=0.9,
        status=ProposalStatus.PENDING,
    )
    db_session.add(proposal)
    await db_session.flush()

    await decide(db_session, proposal.id, actor=actor, decision=Decision.ACCEPT)

    draft = await db_session.scalar(select(AnswerDraft))
    assert draft.body == payload["body"]
    assert draft.cited_clause_ids == [clause.id]
    assert question.status is QuestionStatus.DRAFTED


@pytest.mark.asyncio
async def test_finalize_moves_answer_into_history(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (await import_questions(db_session, engagement.id, "Question?", "en", actor=actor))[
        0
    ]
    answer = AnswerDraft(
        question_id=question.id,
        body="Final answer",
        language="en",
        cited_clause_ids=[],
        cited_control_ids=[],
        suggested_evidence_ids=[],
        gap_notes="",
        reviewed_by=actor.id,
    )
    db_session.add(answer)
    await db_session.flush()

    await finalize_answer(db_session, answer, actor=actor)
    response = await client.get("/api/audit/history", headers=_headers(actor))

    assert response.status_code == 200
    assert response.json()[0]["question_text"] == "Question?"
    assert response.json()[0]["final_body"] == "Final answer"
    assert question.status is QuestionStatus.FINALIZED


@pytest.mark.asyncio
async def test_viewer_cannot_create_engagement(client, db_session):
    viewer = await _user(db_session, Role.VIEWER)
    response = await client.post(
        "/api/audit/engagements",
        headers=_headers(viewer),
        json={"name": "Audit", "audit_type": "internal", "status": "preparing"},
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_generation_creates_answer_proposal_not_formal_draft(db_session, monkeypatch):
    from app.audit import service

    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (
        await import_questions(
            db_session, engagement.id, "How is access reviewed?", "en", actor=actor
        )
    )[0]
    document = Document(
        title="Policy",
        doc_type=DocType.POLICY,
        file_hash="b" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number="1",
        heading="Review",
        heading_path="Review",
        citation_label="1",
        text="Access is reviewed quarterly.",
        order_index=0,
        level=1,
    )
    db_session.add(clause)
    await db_session.flush()

    async def fake_search(*args, **kwargs):
        return SearchResponse(
            query=question.question_text,
            expanded_terms=[],
            vector_used=False,
            hits=[
                SearchHit(
                    chunk_id=1,
                    clause_id=clause.id,
                    document_id=document.id,
                    document_title=document.title,
                    citation_label="1",
                    heading_path="Review",
                    text=clause.text,
                    score=1,
                    rank_fulltext=1,
                    rank_vector=None,
                )
            ],
        )

    async def fake_run(*args, **kwargs):
        return ValidatedResult(
            payload={
                "body": "Access is reviewed quarterly.",
                "language": "en",
                "citations": [{"clause_id": clause.id, "quote": "Access is reviewed quarterly."}],
                "cited_control_ids": [],
                "suggested_evidence_ids": [],
                "gap_notes": "",
                "confidence": 0.9,
            },
            confidence=0.9,
            llm_call_id=None,
        )

    monkeypatch.setattr(service, "search", fake_search)
    monkeypatch.setattr(service, "run", fake_run)
    proposal = await service.generate_answer(db_session, question.id, "en", actor=actor)

    assert proposal.kind is ProposalKind.ANSWER
    assert proposal.status is ProposalStatus.PENDING
    assert await db_session.scalar(select(AnswerDraft)) is None


@pytest.mark.asyncio
async def test_answer_validator_rejects_control_outside_retrieval(db_session):
    document = Document(
        title="Policy",
        doc_type=DocType.POLICY,
        file_hash="c" * 64,
        file_path="/p.pdf",
        original_filename="p.pdf",
        status=DocStatus.ACTIVE,
    )
    db_session.add(document)
    await db_session.flush()
    clause = Clause(
        document_id=document.id,
        number="1",
        heading="Review",
        heading_path="Review",
        citation_label="1",
        text="Access is reviewed quarterly.",
        order_index=0,
        level=1,
    )
    db_session.add(clause)
    await db_session.flush()
    validator = AnswerCitationValidator(
        db_session,
        allowed_clause_ids={clause.id},
        allowed_control_ids=set(),
        allowed_evidence_ids=set(),
    )

    reason = await validator.check(
        {
            "citations": [{"clause_id": clause.id, "quote": clause.text}],
            "cited_control_ids": [999],
            "suggested_evidence_ids": [],
        }
    )

    assert reason == "控制项 999 不属于本次检索结果"


@pytest.mark.asyncio
async def test_generation_rejects_a_second_pending_answer(db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    question = (await import_questions(db_session, engagement.id, "Question?", "en", actor=actor))[
        0
    ]
    db_session.add(
        Proposal(
            kind=ProposalKind.ANSWER,
            payload={"question_id": question.id},
            citations=[],
            confidence=0.8,
            status=ProposalStatus.PENDING,
        )
    )
    await db_session.flush()

    with pytest.raises(Conflict, match="待审核"):
        await generate_answer(db_session, question.id, "en", actor=actor)


@pytest.mark.asyncio
async def test_scoped_viewer_only_sees_its_engagement(client, db_session):
    lead = await _user(db_session)
    allowed = await _engagement(db_session, lead)
    hidden = AuditEngagement(
        name="Hidden audit",
        audit_type=AuditType.INTERNAL,
        status=EngagementStatus.PREPARING,
        created_by=lead.id,
    )
    db_session.add(hidden)
    await db_session.flush()
    viewer = await _user(db_session, Role.VIEWER)
    viewer.engagement_scope_id = allowed.id
    await db_session.flush()

    visible = await client.get("/api/audit/engagements", headers=_headers(viewer))
    blocked = await client.get(
        f"/api/audit/engagements/{hidden.id}/questions", headers=_headers(viewer)
    )

    assert [row["id"] for row in visible.json()] == [allowed.id]
    assert blocked.status_code == 404


def _question_workbook(rows: list[list[str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["审计问题", "语言"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


@pytest.mark.asyncio
async def test_excel_import_preserves_each_row_language(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)

    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/questions/xlsx",
        headers=_headers(actor),
        files={
            "file": (
                "questions.xlsx",
                _question_workbook(
                    [
                        ["How is privileged access approved?", "English"],
                        ["如何复核特权账号？", "中文"],
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    assert [(row["question_text"], row["language"]) for row in response.json()] == [
        ("How is privileged access approved?", "en"),
        ("如何复核特权账号？", "zh"),
    ]
    audit = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_questions.import")
    )
    assert audit.after["source"] == "xlsx"


@pytest.mark.asyncio
async def test_word_export_contains_finalized_answers_only(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    finalized, pending = await import_questions(
        db_session, engagement.id, "Final question?\nPending question?", "en", actor=actor
    )
    answer = AnswerDraft(
        question_id=finalized.id,
        body="Approved response.",
        language="en",
        cited_clause_ids=[],
        cited_control_ids=[],
        suggested_evidence_ids=[],
        gap_notes="",
        reviewed_by=actor.id,
    )
    db_session.add(answer)
    await db_session.flush()
    await finalize_answer(db_session, answer, actor=actor)

    response = await client.get(
        f"/api/audit/engagements/{engagement.id}/export.docx?language=en",
        headers=_headers(actor),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = WordDocument(BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Final question?" in text
    assert "Approved response." in text
    assert pending.question_text not in text
    assert await db_session.scalar(
        select(AuditLog.id).where(AuditLog.action == "audit_engagement.export")
    )


def test_question_similarity_supports_english_and_chinese():
    assert similarity(
        "How is privileged access reviewed?",
        "How do you review privileged access accounts?",
    ) > similarity(
        "How is privileged access reviewed?",
        "How are backups restored?",
    )
    assert similarity("如何复核特权账号？", "特权账号如何定期复核？") > 0.3


@pytest.mark.asyncio
async def test_similar_history_returns_closest_finalized_answer(client, db_session):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    old_question, unrelated, current = await import_questions(
        db_session,
        engagement.id,
        "How is privileged access reviewed?\nHow are backups restored?\nHow do you review privileged access accounts?",
        "en",
        actor=actor,
    )
    for question, body in ((old_question, "Quarterly."), (unrelated, "From snapshots.")):
        answer = AnswerDraft(
            question_id=question.id,
            body=body,
            language="en",
            cited_clause_ids=[],
            cited_control_ids=[],
            suggested_evidence_ids=[],
            gap_notes="",
            reviewed_by=actor.id,
        )
        db_session.add(answer)
        await db_session.flush()
        await finalize_answer(db_session, answer, actor=actor)

    rows = await similar_history(db_session, current.question_text, exclude_question_id=current.id)
    response = await client.get(
        f"/api/audit/questions/{current.id}/similar-history", headers=_headers(actor)
    )

    assert rows[0].question_id == old_question.id
    assert response.status_code == 200
    assert response.json()[0]["answer"] == "Quarterly."
    assert response.json()[0]["similarity"] > 0


@pytest.mark.asyncio
async def test_batch_generation_enqueues_pending_questions(client, db_session, monkeypatch):
    actor = await _user(db_session)
    engagement = await _engagement(db_session, actor)
    await import_questions(
        db_session, engagement.id, "Question one?\nQuestion two?", "en", actor=actor
    )
    calls = []

    async def fake_enqueue(function, *args):
        calls.append((function, args))
        return "audit-job-1"

    monkeypatch.setattr("app.audit.router.enqueue", fake_enqueue)
    response = await client.post(
        f"/api/audit/engagements/{engagement.id}/generate-all",
        headers=_headers(actor),
    )

    assert response.status_code == 200
    assert response.json() == {"job_id": "audit-job-1", "question_count": 2}
    assert calls == [("generate_engagement_answers", (engagement.id, actor.id))]
    audit = await db_session.scalar(
        select(AuditLog).where(AuditLog.action == "audit_answers.batch_enqueue")
    )
    assert audit.after["question_count"] == 2
