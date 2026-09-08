from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document as DocxDocument
from sqlalchemy import select

from app.clauses.models import Clause
from app.ingest.models import DocStatus, DocType, Document
from app.ingest.tasks import run_parse


def _make_docx(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_paragraph("Version: \t1.0")
    doc.add_paragraph("Owner: \tIT Department")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Purpose text.")
    doc.add_heading("Roles and Responsibilities", level=1)
    doc.add_paragraph("IT Division drafts it.")
    path = tmp_path / "g.docx"
    doc.save(path)
    return path


async def _doc(db_session, path: Path, **kwargs) -> Document:
    defaults = dict(
        title="Guideline",
        doc_type=DocType.GUIDELINE,
        file_hash="f" * 64,
        file_path=str(path),
        original_filename=path.name,
    )
    defaults.update(kwargs)
    document = Document(**defaults)
    db_session.add(document)
    await db_session.flush()
    return document


@pytest.mark.asyncio
async def test_successful_parse_activates_document(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    result = await run_parse(db_session, document)

    assert document.status is DocStatus.ACTIVE
    assert result["clauses"] >= 2
    assert document.parse_error is None


@pytest.mark.asyncio
async def test_worker_moves_root_tmp_upload_into_document_store(db_session, tmp_path):
    source = tmp_path / "staged.docx"
    _make_docx(tmp_path).replace(source)
    document = await _doc(db_session, source)
    document.file_path = f"/tmp/{source.name}"
    source.rename(document.file_path)

    from app.parsing.contract import DocumentMeta, ParsedDocument

    stored_path = tmp_path / "stored.docx"
    parsed = ParsedDocument(
        meta=DocumentMeta(
            title=None,
            version=None,
            owner=None,
            approver=None,
            approved_date=None,
            effective_date=None,
            doc_type=None,
        ),
        clauses=[],
    )
    with patch("app.ingest.tasks.save") as save_mock, patch(
        "app.ingest.tasks.get_parser"
    ) as get_parser_mock:
        save_mock.return_value.path = str(stored_path)
        get_parser_mock.return_value.parse.return_value = parsed
        await run_parse(db_session, document)

    save_mock.assert_called_once()
    assert document.file_path == str(stored_path)


@pytest.mark.asyncio
async def test_parse_writes_clauses(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    await run_parse(db_session, document)

    headings = [clause.heading for clause in await db_session.scalars(select(Clause))]
    assert "Introduction" in headings
    assert "Roles and Responsibilities" in headings


@pytest.mark.asyncio
async def test_cover_metadata_is_written_back(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    await run_parse(db_session, document)

    assert document.version == "1.0"
    assert document.owner == "IT Department"


@pytest.mark.asyncio
async def test_parse_failure_records_a_specific_reason(db_session, tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx")
    document = await _doc(db_session, broken)

    with pytest.raises(Exception):
        await run_parse(db_session, document)

    assert document.status is DocStatus.PARSE_FAILED
    assert document.parse_error
    assert "docx" in document.parse_error.lower()


@pytest.mark.asyncio
async def test_completeness_warning_is_stored(db_session, tmp_path):
    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("x")
    path = tmp_path / "partial.docx"
    doc.save(path)

    document = await _doc(db_session, path)
    result = await run_parse(db_session, document)

    assert any("Roles and Responsibilities" in warning for warning in result["warnings"])
    assert "Roles and Responsibilities" in (document.parse_warnings or "")


@pytest.mark.asyncio
async def test_reparse_does_not_duplicate_clauses(db_session, tmp_path):
    document = await _doc(db_session, _make_docx(tmp_path))
    first = await run_parse(db_session, document)
    second = await run_parse(db_session, document)
    assert first["clauses"] == second["clauses"]

    rows = list(await db_session.scalars(select(Clause)))
    assert len(rows) == first["clauses"]


@pytest.mark.asyncio
async def test_low_confidence_ocr_sets_the_quality_flag(db_session, tmp_path):
    from app.parsing.contract import DocumentMeta, ParsedDocument

    document = await _doc(db_session, _make_docx(tmp_path))
    parsed = ParsedDocument(
        meta=DocumentMeta(
            title=None,
            version=None,
            owner=None,
            approver=None,
            approved_date=None,
            effective_date=None,
            doc_type=None,
        ),
        clauses=[],
        warnings=["OCR 置信度偏低"],
        ocr_used=True,
    )
    with patch("app.ingest.tasks.get_parser") as get_parser_mock:
        get_parser_mock.return_value.parse.return_value = parsed
        await run_parse(db_session, document)

    assert document.ocr_quality_flag is True
