from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.parsing.contract import ParseError
from app.parsing.docx_parser import DocxParser


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_paragraph("INTERNAL")
    doc.add_paragraph("Acme Bank")
    doc.add_paragraph("IT Guideline")
    doc.add_paragraph("IT Incident Management")
    doc.add_paragraph("Version: \t\t1.0")
    doc.add_paragraph("Owner: \t\tIT Department")
    doc.add_paragraph("Approver: \t\tChief Risk Officer")
    doc.add_paragraph("Effective Date: \t01/06/2025")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Definition"
    table.cell(1, 0).text = "IT Incident"
    table.cell(1, 1).text = "An unplanned interruption to an IT service."

    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("The purpose of this Guideline is to define incident handling.")
    doc.add_heading("Purpose and Objective", level=2)
    doc.add_paragraph("Defines objectives.")
    doc.add_heading("", level=2)
    doc.add_heading("Periodic Review", level=2)
    doc.add_paragraph("Reviewed annually.")
    doc.add_heading("Roles and Responsibilities", level=1)
    doc.add_heading("IT Department", level=2)
    doc.add_paragraph("Responsible for drafting.")

    path = tmp_path / "guideline.docx"
    doc.save(path)
    return path


def _walk(nodes):
    for node in nodes:
        yield node
        yield from _walk(node.children)


def test_headings_have_no_numbers(sample_docx: Path):
    parsed = DocxParser().parse(sample_docx)
    sections = [node for node in parsed.clauses if node.kind == "section"]
    assert sections
    assert all(node.number is None for node in sections)


def test_hierarchy_comes_from_heading_styles(sample_docx: Path):
    parsed = DocxParser().parse(sample_docx)
    tops = [node for node in parsed.clauses if node.kind == "section"]

    assert [node.heading for node in tops] == ["Introduction", "Roles and Responsibilities"]
    assert [child.heading for child in tops[0].children] == [
        "Purpose and Objective",
        "Periodic Review",
    ]
    assert [child.heading for child in tops[1].children] == ["IT Department"]


def test_empty_headings_do_not_become_clauses(sample_docx: Path):
    parsed = DocxParser().parse(sample_docx)
    assert all(node.heading.strip() for node in _walk(parsed.clauses))


def test_body_text_is_attached_to_its_heading(sample_docx: Path):
    parsed = DocxParser().parse(sample_docx)
    intro = next(node for node in parsed.clauses if node.heading == "Introduction")
    assert "define incident handling" in intro.text
    assert "Reviewed annually" in intro.children[1].text


def test_tables_become_clauses(sample_docx: Path):
    parsed = DocxParser().parse(sample_docx)
    tables = [node for node in parsed.clauses if node.kind == "table"]

    assert len(tables) == 1
    assert "IT Incident" in tables[0].text
    assert "unplanned interruption" in tables[0].text


def test_cover_metadata_is_extracted(sample_docx: Path):
    from datetime import date

    meta = DocxParser().parse(sample_docx).meta
    assert meta.version == "1.0"
    assert meta.owner == "IT Department"
    assert meta.approver == "Chief Risk Officer"
    assert meta.effective_date == date(2025, 6, 1)


def test_doc_type_is_inferred_from_cover(sample_docx: Path):
    assert DocxParser().parse(sample_docx).meta.doc_type == "guideline"


def test_corrupt_file_raises_a_readable_parse_error(tmp_path: Path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx at all")

    with pytest.raises(ParseError) as exc:
        DocxParser().parse(broken)
    assert "docx" in exc.value.reason.lower()


def test_document_without_headings_warns_instead_of_silently_returning_nothing(tmp_path: Path):
    doc = DocxDocument()
    doc.add_paragraph("just a flat paragraph")
    path = tmp_path / "flat.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)
    assert parsed.clauses == [] or all(node.kind == "table" for node in parsed.clauses)
    assert any("标题" in warning for warning in parsed.warnings)


def _tracked_paragraph(doc, before: str, inserted: str, after: str, deleted: str = ""):
    """造一个含修订痕迹的段落：inserted 在 w:ins 里，deleted 在 w:del 里。"""
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph(before)
    p = paragraph._element

    ins = p.makeelement(qn("w:ins"), {qn("w:id"): "1", qn("w:author"): "a"})
    run = ins.makeelement(qn("w:r"), {})
    text = run.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
    text.text = inserted
    run.append(text)
    ins.append(run)
    p.append(ins)

    if deleted:
        dele = p.makeelement(qn("w:del"), {qn("w:id"): "2", qn("w:author"): "a"})
        drun = dele.makeelement(qn("w:r"), {})
        dtext = drun.makeelement(qn("w:delText"), {qn("xml:space"): "preserve"})
        dtext.text = deleted
        drun.append(dtext)
        dele.append(drun)
        p.append(dele)

    tail = paragraph.add_run(after)
    tail.text = after
    return paragraph


@pytest.fixture
def tracked_changes_docx(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_heading("Roles and Responsibilities", level=1)
    _tracked_paragraph(
        doc,
        before="If the ",
        inserted="Incident Manager",
        after=" has contacted the requestor three times, the ticket may be closed.",
        deleted="OBSOLETE WORDING",
    )
    table = doc.add_table(rows=1, cols=1)
    cell_paragraph = table.cell(0, 0).paragraphs[0]
    from docx.oxml.ns import qn

    ins = cell_paragraph._element.makeelement(qn("w:ins"), {qn("w:id"): "3", qn("w:author"): "a"})
    run = ins.makeelement(qn("w:r"), {})
    text = run.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
    text.text = "Chief Operating Officer"
    run.append(text)
    ins.append(run)
    cell_paragraph._element.append(ins)

    path = tmp_path / "tracked.docx"
    doc.save(path)
    return path


def test_tracked_insertions_are_kept_and_deletions_dropped(tracked_changes_docx: Path):
    """python-docx 的 Paragraph.text 只看 w:p 的直接 w:r 子节点。

    嵌在 w:ins 里的 run 会被静默丢弃——实测一份带修订的真实文档丢了 96%
    的插入片段，句子中间凭空少词。下游没有任何闸门能察觉：引文确实逐字
    存在于（残缺的）条款正文里，模型却会自行补上一个看似合理的主语。
    """
    parsed = DocxParser().parse(tracked_changes_docx)
    body = "\n".join(node.text for node in parsed.clauses) + "\n".join(
        child.text for node in parsed.clauses for child in node.children
    )

    assert "Incident Manager" in body, "修订插入的文字被丢了"
    assert "If the Incident Manager  has contacted" in body.replace("\n", " ") or \
           "If the Incident Manager has contacted" in body.replace("  ", " ")
    assert "OBSOLETE WORDING" not in body, "修订删除的文字不该出现"


def test_tracked_insertions_inside_tables_are_kept(tracked_changes_docx: Path):
    parsed = DocxParser().parse(tracked_changes_docx)
    everything = str(parsed.clauses) + "".join(
        node.text for node in parsed.clauses
    ) + "".join(child.text for node in parsed.clauses for child in node.children)
    assert "Chief Operating Officer" in everything


# ── 没有 Heading 样式时，从编号认标题 ──────────────────────────
# 实测 14 份外部机构的 .docx，段落样式 **100% 是 `Normal`**（154/154、42/42、
# 1218/1218）：视觉上的标题是加粗和字号做出来的，不套样式。而标题识别原本
# 只认样式，于是一条章节都找不到，产出的每一条"条款"都是被单独捞出来的表格。
# 样式仍然优先——样本语料靠它，且它比任何启发式都可靠。


def _normal_docx(tmp_path: Path, lines: list[str], name: str = "normal.docx") -> Path:
    """全部用 Normal 样式写一份文档，标题只靠编号区分。"""
    doc = DocxDocument()
    for line in lines:
        doc.add_paragraph(line)
    path = tmp_path / name
    doc.save(path)
    return path


def test_a_document_without_heading_styles_gets_its_tree_from_the_numbering(tmp_path: Path):
    path = _normal_docx(tmp_path, [
        "1. Introduction",
        "This standard applies to all systems.",
        "1.1. Purpose",
        "To set out the requirements.",
        "1.2. Scope",
        "All production environments.",
        "2. Requirements",
        "The following apply.",
    ])

    parsed = DocxParser().parse(path)

    numbers = [node.number for node in parsed.clauses]
    assert numbers == ["1", "2"]
    assert [child.number for child in parsed.clauses[0].children] == ["1.1", "1.2"]
    assert "applies to all systems" in parsed.clauses[0].text


def test_body_list_items_do_not_become_sections(tmp_path: Path):
    """BCBS 的写法：`1.` `2.` `3.` 是正文列表，没有任何子号挂上去。"""
    # 关键是**重复**：正文列表在每一节里各起一遍，扁平章节编号只升一次。
    path = _normal_docx(tmp_path, [
        "Principles",
        "1. Boards should approve the strategy.",
        "2. Boards should oversee implementation.",
        "3. Boards should review outcomes.",
        "Responsibilities",
        "1. Management should implement the strategy.",
        "2. Management should report on it.",
    ])

    parsed = DocxParser().parse(path)

    assert [node.number for node in parsed.clauses if node.kind == "section"] == []


def test_heading_styles_still_win_when_the_document_has_them(tmp_path: Path):
    """样本语料靠样式。有样式就不猜编号。"""
    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("1. This numbered line is body text, not a section.")
    path = tmp_path / "styled.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)

    assert [node.heading for node in parsed.clauses] == ["Introduction"]
    assert "1. This numbered line" in parsed.clauses[0].text


def test_a_table_lands_under_the_numbered_section_above_it(tmp_path: Path):
    """Task 4 与 Task 5 合起来才完整：先有章节，表格才有地方挂。"""
    doc = DocxDocument()
    doc.add_paragraph("4.1. End-User Passwords")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Length"
    table.cell(0, 1).text = "At least twelve characters"
    doc.add_paragraph("4.2. Privileged Passwords")
    path = tmp_path / "numbered-with-table.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)

    # `4` 从未出现，补出来当父节点；表格挂在真正的 `4.1` 下面
    assert parsed.clauses[0].number == "4"
    first = parsed.clauses[0].children[0]
    assert first.number == "4.1"
    assert [child.kind for child in first.children] == ["table"]


def test_a_numbered_child_hangs_under_its_own_parent_not_whatever_came_before(
    tmp_path: Path,
):
    """按层级入栈是不够的：层级对得上，编号不一定对得上。

    实测 01_Arab_Bank / 02_Bank_of_Jordan / 06_OCC / 10_RCBC 四份都栽在这里——
    条款树长出来了，但 `4.1` 挂在了 `3` 下面。审计引用顺着它走会走到别的章节去。
    """
    doc = DocxDocument()
    # 关键是 `4.1` 出现时 `4` 从未出现过：按层级入栈会把它挂到上一个一级节点
    # `3` 下面去，层级对得上、编号对不上。
    for line in ["3. Governance", "3.1. Roles", "4.1. Passwords", "4.2. Tokens"]:
        doc.add_paragraph(line)
    path = tmp_path / "lineage.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)

    def check(parent):
        for child in parent.children:
            if parent.number and child.number:
                assert child.number.startswith(f"{parent.number}."), (
                    f"{child.number} 挂在了 {parent.number} 下面"
                )
            check(child)

    assert [node.number for node in parsed.clauses] == ["3", "4"]
    for root in parsed.clauses:
        check(root)


# ── 既无样式也无编号时，从格式认标题 ──────────────────────────
# 剩下的委员会章程类文档既不套 Heading 样式，也不给章节编号——标题是
# "整段加粗" 或 "字号比正文大" 做出来的。实测候选数都很合理：
# 03 有 4 个 / 47 段，07 有 7 / 28，13 有 5 / 42，15 有 29 / 336。
#
# 判据一律是**相对**的：比这份文档的正文字号大、整段加粗且短。
# 不写死 12pt / 14pt——那又会变成一份文档的特征。


def _formatted_docx(tmp_path: Path, blocks: list[tuple[str, bool]], name="fmt.docx") -> Path:
    """blocks 里每项是 (文本, 是否整段加粗)。全部 Normal 样式，无编号。"""
    doc = DocxDocument()
    for text, bold in blocks:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = bold
    path = tmp_path / name
    doc.save(path)
    return path


def test_bold_short_lines_become_sections_when_nothing_else_marks_them(tmp_path: Path):
    path = _formatted_docx(tmp_path, [
        ("Purpose", True),
        ("The Committee assists the Board in overseeing technology.", False),
        ("Composition", True),
        ("The Committee shall consist of at least three directors.", False),
    ])

    parsed = DocxParser().parse(path)

    assert [node.heading for node in parsed.clauses] == ["Purpose", "Composition"]
    assert "assists the Board" in parsed.clauses[0].text
    assert "at least three directors" in parsed.clauses[1].text


def test_a_long_bold_paragraph_is_not_a_heading(tmp_path: Path):
    """整段加粗的强调段落不是标题——标题短。"""
    path = _formatted_docx(tmp_path, [
        ("Purpose", True),
        ("The Committee assists the Board in overseeing technology strategy, "
         "cyber risk, data governance and the technology investment portfolio, "
         "and reports to the Board after each meeting.", True),
    ])

    parsed = DocxParser().parse(path)

    assert [node.heading for node in parsed.clauses] == ["Purpose"]


def test_numbering_wins_over_formatting(tmp_path: Path):
    """有编号就用编号——它能给出层级，格式给不出。"""
    doc = DocxDocument()
    for text, bold in [("1. Purpose", False), ("Some bold emphasis", True),
                       ("2. Composition", False), ("3. Reporting", False)]:
        paragraph = doc.add_paragraph()
        paragraph.add_run(text).bold = bold
    path = tmp_path / "numbered-wins.docx"
    doc.save(path)

    parsed = DocxParser().parse(path)

    assert [node.number for node in parsed.clauses] == ["1", "2", "3"]


def test_formatting_headings_carry_no_number(tmp_path: Path):
    """认不出编号就别编一个——citation_label 会退化成标题路径，那是诚实的。"""
    path = _formatted_docx(tmp_path, [("Reporting", True), ("Body text here.", False)])

    parsed = DocxParser().parse(path)

    assert parsed.clauses[0].number is None
