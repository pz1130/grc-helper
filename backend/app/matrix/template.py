"""Bounded XLSX parsing and pure validation; never evaluates spreadsheet formulas."""

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

CANONICAL_FIELDS: tuple[str, ...] = (
    "code", "title", "statement", "category", "owner", "framework_refs", "note",
)
REQUIRED_FIELDS: tuple[str, ...] = ("title", "statement")
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 128
MAX_CELL_CHARS = 32_767
MAX_TEXT_CHARS = 8 * 1024 * 1024


@dataclass(frozen=True)
class Sheet:
    headers: list[str]
    rows: list[list[str]]


def _inspect_archive(content: bytes) -> None:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Excel 文件为空或超过 10 MiB 限制")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 256 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
            raise ValueError("Excel 解压大小或文件数量超过限制")
        if len({e.filename for e in entries}) != len(entries):
            raise ValueError("Excel 包含重复文件")
        for entry in entries:
            if entry.flag_bits & 1 or entry.filename.lower().endswith(".bin"):
                raise ValueError("Excel 不支持加密文件或宏")
            if not entry.filename.endswith((".xml", ".rels")):
                continue
            raw = archive.read(entry)
            # XLSX XML is UTF-8. Reject DTD/entity declarations, including UTF-16
            # spellings, before either XML parser sees untrusted declarations.
            safe_scan = raw.replace(b"\x00", b"").upper()
            if b"<!DOCTYPE" in safe_scan or b"<!ENTITY" in safe_scan:
                raise ValueError("Excel XML 不允许实体或 DTD")
            if not entry.filename.startswith("xl/worksheets/"):
                continue
            for _, element in ElementTree.iterparse(io.BytesIO(raw), events=("start",)):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "row" and int(element.get("r", "1")) > MAX_ROWS + 1:
                    raise ValueError("Excel 行数超过限制")
                if tag == "c":
                    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", element.get("r", ""))
                    if not match:
                        raise ValueError("Excel 单元格坐标无效")
                    if (column_index_from_string(match[1]) > MAX_COLUMNS
                            or int(match[2]) > MAX_ROWS + 1):
                        raise ValueError("Excel 行列数超过限制")
                element.clear()


def read_sheet(content: bytes, *, max_rows: int | None = None) -> Sheet:
    if max_rows is not None and (type(max_rows) is not int or max_rows < 0):
        raise ValueError("max_rows 必须为非负整数")
    book = None
    try:
        _inspect_archive(content)
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        worksheet = book.active
        if worksheet is None or not hasattr(worksheet, "iter_rows"):
            raise ValueError("Excel 缺少数据工作表")
        # Do not trust attacker-controlled worksheet dimensions (e.g. XFD1048576).
        worksheet.reset_dimensions()
        headers: list[str] = []
        rows: list[list[str]] = []
        text_chars = 0
        for number, cells in enumerate(worksheet.iter_rows(), start=1):
            if number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                raise ValueError("Excel 行列数超过限制")
            row: list[str] = []
            for cell in cells:
                if cell.data_type in {"f", "e"}:
                    raise ValueError("Excel 包含公式或错误单元格，请先粘贴为值")
                value = "" if cell.value is None else str(cell.value).strip()
                if len(value) > MAX_CELL_CHARS:
                    raise ValueError("Excel 单元格文本超过限制")
                text_chars += len(value)
                if text_chars > MAX_TEXT_CHARS:
                    raise ValueError("Excel 文本总量超过限制")
                row.append(value)
            if number == 1:
                headers = row
                while headers and not headers[-1]:
                    headers.pop()
                if max_rows == 0:
                    break
            elif any(row):
                rows.append(row)
                if max_rows is not None and len(rows) >= max_rows:
                    break
        return Sheet(headers, rows)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"无法读取 Excel 文件：{type(exc).__name__}") from exc
    finally:
        if book is not None:
            book.close()


def validate(sheet: Sheet, mapping: dict[str, str]) -> list[str]:
    report: list[str] = []
    if not isinstance(mapping, dict) or not all(
        isinstance(k, str) and isinstance(v, str) for k, v in mapping.items()
    ):
        return ["mapping 必须为目标字段到源列名的字符串对象"]
    if not sheet.headers or any(not h for h in sheet.headers):
        report.append("Excel 表头为空或包含空列名")
    if len(set(sheet.headers)) != len(sheet.headers):
        report.append("Excel 表头重复，无法确定源列")
    for field in REQUIRED_FIELDS:
        if field not in mapping:
            report.append(f"必填字段 {field} 没有映射到任何一列")
    index: dict[str, int] = {}
    for field, column in mapping.items():
        if field not in CANONICAL_FIELDS:
            report.append(f"未知的目标字段 {field}")
        elif column not in sheet.headers:
            report.append(f"映射到了不存在的列「{column}」")
        else:
            index[field] = sheet.headers.index(column)
    if report:
        return report
    if not sheet.rows:
        report.append("Excel 没有数据行")
    for field in REQUIRED_FIELDS:
        blanks = sum(not _cell(row, index[field]) for row in sheet.rows)
        if blanks:
            report.append(f"有 {blanks} 行的 {field} 为空")
    for field, limit in (("code", 64), ("title", 500), ("category", 100)):
        if field in index and any(len(_cell(r, index[field])) > limit for r in sheet.rows):
            report.append(f"{field} 长度超过 {limit}")
    if any(any(row[len(sheet.headers):]) for row in sheet.rows):
        report.append("存在没有表头的数据列")
    if "code" in index:
        counts = Counter(_cell(row, index["code"]) for row in sheet.rows)
        duplicates = sorted(code for code, count in counts.items() if code and count > 1)
        if duplicates:
            report.append(f"编号重复：{'、'.join(duplicates[:20])}")
    return report


def _cell(row: list[str], position: int) -> str:
    return row[position].strip() if position < len(row) else ""
