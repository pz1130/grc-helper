"""Create a client-ready Word package from finalized audit answers."""

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.audit.models import AnswerDraft, AuditEngagement, AuditQuestion, QuestionStatus
from app.clauses.models import Clause
from app.controls.models import Control
from app.errors import AppError
from app.evidence.models import EvidenceItem
from app.iam.models import User
from app.ingest.models import Document
from app.review.models import Proposal, ProposalKind, ProposalStatus


@dataclass(frozen=True)
class ExportRow:
    question: AuditQuestion
    answer: AnswerDraft
    citations: list[dict]


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    properties.append(borders)


def _set_cell_margin(cell, value: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _labels(language: str) -> dict[str, str]:
    if language == "zh":
        return {
            "title": "审计答复包",
            "type": "审计类型",
            "period": "审计期间",
            "generated": "导出时间",
            "count": "已定稿答复",
            "question": "问题",
            "answer": "答复",
            "references": "内部依据",
            "evidence": "证据清单",
            "gaps": "差距与后续事项",
            "none": "无",
            "items": "项",
            "e_title": "证据",
            "control": "控制点",
            "status": "状态",
            "owner": "责任人",
            "valid": "有效期至",
            "location": "存放位置",
            "details": "详情",
        }
    return {
        "title": "Audit Response Package",
        "type": "Audit type",
        "period": "Audit period",
        "generated": "Exported",
        "count": "Finalized answers",
        "question": "Question",
        "answer": "Answer",
        "references": "Internal references",
        "evidence": "Evidence list",
        "gaps": "Gaps and follow-up",
        "none": "None",
        "items": "items",
        "e_title": "Evidence",
        "control": "Control",
        "status": "Status",
        "owner": "Owner",
        "valid": "Valid until",
        "location": "Location",
        "details": "Details",
    }


def _configure(document: WordDocument) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    styles = document.styles
    for name, size in (("Normal", 11), ("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)


def _period(engagement: AuditEngagement) -> str:
    if engagement.period_start and engagement.period_end:
        return f"{engagement.period_start.isoformat()} to {engagement.period_end.isoformat()}"
    return (
        engagement.period_start.isoformat()
        if engagement.period_start
        else (engagement.period_end.isoformat() if engagement.period_end else "—")
    )


async def build_docx(session: AsyncSession, engagement: AuditEngagement, *, language: str) -> bytes:
    labels = _labels(language)
    answer_rows = list(
        await session.execute(
            select(AuditQuestion, AnswerDraft)
            .join(AnswerDraft, AnswerDraft.question_id == AuditQuestion.id)
            .where(
                AuditQuestion.engagement_id == engagement.id,
                AuditQuestion.status == QuestionStatus.FINALIZED,
                AnswerDraft.finalized_at.is_not(None),
            )
            .order_by(AuditQuestion.seq)
        )
    )
    if not answer_rows:
        raise AppError("该审计项目还没有已定稿答复")
    question_ids = {question.id for question, _ in answer_rows}
    proposals = list(
        await session.scalars(
            select(Proposal)
            .where(
                Proposal.kind == ProposalKind.ANSWER,
                Proposal.status.in_((ProposalStatus.ACCEPTED, ProposalStatus.MODIFIED)),
            )
            .order_by(Proposal.id)
        )
    )
    proposal_payloads = {
        payload["question_id"]: payload
        for proposal in proposals
        if (payload := proposal.decided_payload or proposal.payload).get("question_id")
        in question_ids
    }
    rows = [
        ExportRow(question, answer, proposal_payloads.get(question.id, {}).get("citations", []))
        for question, answer in answer_rows
    ]
    clause_ids = {citation["clause_id"] for row in rows for citation in row.citations}
    evidence_ids = {item for row in rows for item in row.answer.suggested_evidence_ids}
    clauses = {
        clause.id: clause
        for clause in await session.scalars(select(Clause).where(Clause.id.in_(clause_ids)))
    }
    documents = {
        document.id: document
        for document in await session.scalars(
            select(Document).where(Document.id.in_({row.document_id for row in clauses.values()}))
        )
    }
    evidence = {
        item.id: item
        for item in await session.scalars(
            select(EvidenceItem).where(EvidenceItem.id.in_(evidence_ids))
        )
    }
    control_ids = {item.control_id for item in evidence.values()}
    controls = {
        control.id: control
        for control in await session.scalars(select(Control).where(Control.id.in_(control_ids)))
    }
    owner_ids = {item.owner_user_id for item in evidence.values() if item.owner_user_id}
    owners = {
        user.id: user.name
        for user in await session.scalars(select(User).where(User.id.in_(owner_ids)))
    }

    document = WordDocument()
    _configure(document)
    document.core_properties.title = f"{engagement.name} {labels['title']}"
    document.core_properties.author = "GRC Helper"
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(f"{engagement.name}\n{labels['title']}")
    document.add_paragraph(
        f"This package contains {len(rows)} finalized audit response(s), internal references, and evidence details."
        if language == "en"
        else f"本文件包含 {len(rows)} 项已定稿审计答复及其内部依据和证据清单。"
    )
    metadata = document.add_table(rows=4, cols=2)
    metadata.autofit = False
    metadata.columns[0].width = Inches(1.55)
    metadata.columns[1].width = Inches(5.25)
    meta_rows = (
        (labels["type"], engagement.audit_type.value),
        (labels["period"], _period(engagement)),
        (labels["generated"], datetime.now(UTC).date().isoformat()),
        (labels["count"], f"{len(rows)} {labels['items']}"),
    )
    for index, (key, value) in enumerate(meta_rows):
        metadata.cell(index, 0).text = key
        metadata.cell(index, 1).text = value
        _shade(metadata.cell(index, 0), "EAF1F8")
        metadata.cell(index, 0).paragraphs[0].runs[0].bold = True
        for cell in metadata.rows[index].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margin(cell)
    _borders(metadata)

    for position, row in enumerate(rows):
        if position:
            document.add_page_break()
        document.add_heading(f"{row.question.seq}. {labels['question']}", level=1)
        question_paragraph = document.add_paragraph()
        question_paragraph.add_run(row.question.question_text).bold = True
        document.add_heading(labels["answer"], level=2)
        document.add_paragraph(row.answer.final_body or row.answer.body)
        document.add_heading(labels["references"], level=2)
        if not row.citations:
            document.add_paragraph(labels["none"])
        for citation in row.citations:
            clause = clauses.get(citation.get("clause_id"))
            if clause is None:
                continue
            source = documents.get(clause.document_id)
            heading = document.add_paragraph()
            heading.add_run(
                f"{source.title if source else 'Document'} — {clause.citation_label}"
            ).bold = True
            quote = document.add_paragraph(citation.get("quote", ""))
            quote.paragraph_format.left_indent = Inches(0.25)
            quote.paragraph_format.right_indent = Inches(0.15)
            quote.runs[0].italic = True
            quote.runs[0].font.color.rgb = RGBColor(70, 70, 70)
        document.add_heading(labels["evidence"], level=2)
        items = [
            evidence[item_id]
            for item_id in row.answer.suggested_evidence_ids
            if item_id in evidence
        ]
        if not items:
            document.add_paragraph(labels["none"])
        else:
            table = document.add_table(rows=1, cols=4)
            table.autofit = False
            widths = (Inches(2.05), Inches(1.9), Inches(0.8), Inches(2.05))
            headers = (labels["e_title"], labels["control"], labels["status"], labels["details"])
            for index, value in enumerate(headers):
                cell = table.cell(0, index)
                cell.text = value
                cell.width = widths[index]
                _shade(cell, "17365D")
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            for row_index, item in enumerate(items, start=1):
                control = controls.get(item.control_id)
                values = (
                    item.title,
                    f"{control.code} {control.title}" if control else str(item.control_id),
                    item.status.value,
                    "\n".join(
                        (
                            f"{labels['owner']}: {owners.get(item.owner_user_id, '—')}",
                            f"{labels['valid']}: {item.valid_until.date().isoformat() if item.valid_until else '—'}",
                            f"{labels['location']}: {item.location_hint or '—'}",
                        )
                    ),
                )
                cells = table.add_row().cells
                for index, value in enumerate(values):
                    cells[index].text = value
                    cells[index].width = widths[index]
                    cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _set_cell_margin(cells[index])
                    if row_index % 2 == 0:
                        _shade(cells[index], "F4F7FA")
            _borders(table)
        if row.answer.gap_notes:
            document.add_heading(labels["gaps"], level=2)
            document.add_paragraph(row.answer.gap_notes)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
