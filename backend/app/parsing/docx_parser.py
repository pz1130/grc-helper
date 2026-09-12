"""DOCX parser based on Heading styles, tables, and cover metadata."""

import re
from datetime import date, datetime
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsing.contract import ClauseNode, DocumentMeta, ParsedDocument, ParseError
from app.parsing.docx_numbering import HeadingNumbering
from app.parsing.headings import acts_as_headings, demote_fragment_headings
from app.parsing.numbering import _noise
from app.parsing.vocab import parse_label

_HEADING = re.compile(r"^Heading (\d+)$")
_META = re.compile(
    r"^\s*(Version|Owner|Approver|Approved Date|Effective Date)\s*:\s*(.+?)\s*$"
)
_DOC_TYPE = re.compile(r"\bIT\s+(Policy|Standard|Procedure|Guideline)\b", re.IGNORECASE)
_DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d %B %Y", "%B %d, %Y")


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_TEXT = f"{_W}t"
_DELETED = f"{_W}del"


def _element_text(element) -> str:
    """收集段落/单元格里的 w:t，跳过 w:del 下的内容。

    不能用 python-docx 的 Paragraph.text：它只拼接 w:p 的**直接** w:r 子节点，
    嵌在 w:ins（修订插入）里的 run 会被静默丢弃。实测一份带修订痕迹的真实
    文档丢了 96% 的插入片段（约 4700 字符），句子中间凭空少词。

    下游没有任何闸门能察觉这件事：引文确实逐字存在于（已残缺的）条款正文里，
    模型却会自行补上一个看似合理的主语，读起来像有据可依。

    w:del 的正文放在 w:delText，本来就取不到；这里显式跳过是防个别写法把
    w:t 塞进 w:del——修订删除的文字不该出现在正文中。
    """
    parts: list[str] = []

    def walk(node) -> None:
        for child in node:
            if child.tag == _DELETED:
                continue
            if child.tag == _TEXT:
                parts.append(child.text or "")
            else:
                walk(child)

    walk(element)
    return "".join(parts)


def _paragraph_text(paragraph: Paragraph) -> str:
    return _element_text(paragraph._element)


def _heading_level(paragraph: Paragraph) -> int | None:
    match = _HEADING.match(paragraph.style.name or "")
    return int(match.group(1)) if match else None


def _parse_date(raw: str) -> date | None:
    for date_format in _DATE_FORMATS:
        try:
            return datetime.strptime(raw.strip(), date_format).date()
        except ValueError:
            continue
    return None


def _table_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [
            "\n".join(_paragraph_text(p) for p in cell.paragraphs).strip()
            for cell in row.cells
        ]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _numbered_heading(text: str, dotted_are_headings: bool):
    """没有 Heading 样式时，从编号认标题。返回 (level, number, title) 或 None。"""
    found = parse_label(text)
    if found is None:
        return None
    if not dotted_are_headings and found.trailing_dot and len(found.parts) == 1:
        return None
    return len(found.parts), ".".join(str(part) for part in found.parts), found.title


# 标题总是短的。80 字符是形状判断：再长就是强调段落，不是标题。
_HEADING_MAX_CHARS = 80


def _body_point_size(paragraphs) -> float | None:
    """这份文档的正文字号——取出现次数最多的那个。"""
    sizes: dict[float, int] = {}
    for paragraph in paragraphs:
        runs = [run for run in paragraph.runs if run.text.strip()]
        points = [run.font.size.pt for run in runs if run.font.size]
        if points:
            sizes[max(points)] = sizes.get(max(points), 0) + 1
    return max(sizes, key=lambda size: sizes[size]) if sizes else None


def _looks_like_heading(paragraph, body_point: float | None) -> bool:
    """既无样式也无编号时的兜底：整段加粗、或字号比正文大，且短。

    判据一律**相对于这份文档**——写死 12pt / 14pt 就又变成一份文档的特征了。
    """
    text = _paragraph_text(paragraph).strip()
    if not text or len(text) > _HEADING_MAX_CHARS:
        return False
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False
    if all(run.bold for run in runs):
        return True
    points = [run.font.size.pt for run in runs if run.font.size]
    return bool(points) and body_point is not None and max(points) > body_point


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise ParseError(f"无法打开 docx 文件：{type(exc).__name__}") from exc

        warnings: list[str] = []
        # Word 的标题编号不在文字里，是渲染时按 numbering.xml 画上去的。
        # 不还原的话 citation_label 只能退化成标题路径，而审计引用要的是 3.4.2。
        numbering = HeadingNumbering(document)
        # 样式优先：样本语料靠它，而且它比任何启发式都可靠。只有当整篇文档
        # **一个 Heading 样式都没有**时才退到编号——实测 14 份外部 docx
        # 全部如此，段落样式 100% 是 Normal。
        styled = any(_heading_level(paragraph) is not None for paragraph in document.paragraphs)
        dotted_are_headings = False
        if not styled:
            survey = [
                found.parts
                for paragraph in document.paragraphs
                if (found := parse_label(_paragraph_text(paragraph).strip()))
            ]
            dotted_are_headings = acts_as_headings(survey)

        # 三级取舍：样式 > 编号 > 格式。样式最可靠，编号能给出层级，
        # 格式只说得出"这是个标题"，给不出层级也给不出编号——所以只在
        # 前两者都无所获时才用它。
        by_format = False
        body_point = None
        if not styled:
            numbered_candidates = sum(
                1
                for paragraph in document.paragraphs
                if _noise(_paragraph_text(paragraph).strip()) is None
                and _numbered_heading(_paragraph_text(paragraph).strip(), dotted_are_headings)
            )
            if numbered_candidates == 0:
                by_format = True
                body_point = _body_point_size(document.paragraphs)
        numbered = 0
        roots: list[ClauseNode] = []
        stack: list[ClauseNode] = []
        meta_fields: dict[str, str] = {}
        doc_type: str | None = None
        empty_headings = 0
        table_index = 0

        def attach(node: ClauseNode, level: int) -> None:
            while stack and stack[-1].level >= level:
                stack.pop()
            if stack:
                stack[-1].children.append(node)
            else:
                roots.append(node)

        # **按文档顺序**遍历 body。`document.paragraphs` 不含表格，
        # `document.tables` 又丢掉位置信息——两个循环各走各的，表格在第几节里
        # 就永远看不出来了。实测一份技术标准因此有六个章节正文为空，
        # 它们的内容是文末孤零零的六张表。
        for element in document.element.body.iterchildren():
            if element.tag == qn("w:tbl"):
                body = _table_text(Table(element, document))
                if not body:
                    continue
                table_index += 1
                node = ClauseNode(
                    heading=f"Table {table_index}",
                    text=body,
                    level=(stack[-1].level + 1) if stack else 1,
                    kind="table",
                )
                if stack:
                    stack[-1].children.append(node)
                else:
                    # 封面上的版本历史表、缩略语表不属于任何章节。
                    roots.append(node)
                continue
            if element.tag != qn("w:p"):
                continue

            paragraph = Paragraph(element, document)
            text = _paragraph_text(paragraph).strip()
            level = _heading_level(paragraph)
            from_number = (
                None
                if styled or _noise(text) is not None
                else _numbered_heading(text, dotted_are_headings)
            )
            if from_number is not None:
                level, number, title = from_number
                numbered += 1
                # 按**编号血统**收栈，不是按层级。层级对得上、编号对不上时，
                # `4.1` 会挂到上一个一级节点 `3` 下面去——审计引用顺着它走
                # 会走到别的章节。实测 01/02/06/10 四份都栽在这里。
                #
                # 缺的祖先要补出来，理由同 numbering.py：不补的话 `4.1` 自己
                # 变成顶层条款，一份文档的"顶层"里混着二级三级编号。
                for depth in range(1, len(number.split("."))):
                    ancestor = ".".join(number.split(".")[:depth])
                    while (
                        stack
                        and stack[-1].number != ancestor
                        and not ancestor.startswith(f"{stack[-1].number}.")
                    ):
                        stack.pop()
                    if stack and stack[-1].number == ancestor:
                        continue
                    placeholder = ClauseNode(
                        heading=ancestor, text="", level=depth, number=ancestor
                    )
                    (stack[-1].children if stack else roots).append(placeholder)
                    stack.append(placeholder)
                while stack and not number.startswith(f"{stack[-1].number}."):
                    stack.pop()
                node = ClauseNode(heading=title, text="", level=level, number=number)
                (stack[-1].children if stack else roots).append(node)
                stack.append(node)
                continue

            if by_format and _looks_like_heading(paragraph, body_point):
                node = ClauseNode(heading=text, text="", level=1, kind="section")
                stack.clear()
                roots.append(node)
                stack.append(node)
                continue

            if level is not None:
                # 计数器必须按文档顺序推进，空标题也要走一遍，否则后面全错位。
                number = numbering.number_for(paragraph)
                if not text:
                    empty_headings += 1
                    continue
                if number:
                    numbered += 1
                node = ClauseNode(heading=text, text="", level=level, number=number)
                attach(node, level)
                stack.append(node)
                continue

            if not text:
                continue

            if match := _META.match(text):
                meta_fields.setdefault(match.group(1), match.group(2).strip())
            if doc_type is None and (match := _DOC_TYPE.search(text)):
                doc_type = match.group(1).lower()

            if stack:
                stack[-1].text = f"{stack[-1].text}\n{text}".strip()

        if empty_headings:
            warnings.append(f"已跳过 {empty_headings} 个空标题")
        if not any(node.kind == "section" for node in roots):
            warnings.append("未找到任何标题样式，该文档可能不是标准 house style")
        if not numbered and any(node.kind == "section" for node in roots):
            # 说出来而不是静默退化：条款号会变成标题路径，审计引用会不好用。
            warnings.append("未能还原任何标题编号，条款号将退化为标题路径")

        meta = DocumentMeta(
            title=None,
            version=meta_fields.get("Version"),
            owner=meta_fields.get("Owner"),
            approver=meta_fields.get("Approver"),
            approved_date=_parse_date(meta_fields.get("Approved Date", "")),
            effective_date=_parse_date(meta_fields.get("Effective Date", "")),
            doc_type=doc_type,
        )
        # 同 PDF 路："标题"其实是段落前半句时，整段收回正文。
        demote_fragment_headings(roots)
        return ParsedDocument(meta=meta, clauses=roots, warnings=warnings)
