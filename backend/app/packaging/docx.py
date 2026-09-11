"""Word member of the audit readiness package."""

from datetime import UTC, datetime
from io import BytesIO

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

from app.audit.export import _borders, _configure, _labels, _shade
from app.conflicts.models import PolicyConflict
from app.frameworks.coverage import CoverageRow, GapRow
from app.frameworks.models import Framework
from app.maturity.schemas import MaturitySummaryOut


def _section_labels(language: str) -> dict[str, str]:
    if language == "zh":
        return {
            "package": "审计准备包",
            "coverage": "覆盖度摘要",
            "gaps": "差距清单",
            "maturity": "成熟度得分",
            "conflicts": "已确认冲突",
            "code": "代码",
            "title": "标题",
            "covered": "已覆盖",
            "requirements": "要求数",
            "supporting": "仅有 supporting",
            "doc": "文档分",
            "impl": "实施分",
            "scored": "已评分",
            "topic": "主题",
            "difference": "差异",
            "yes": "是",
            "no": "否",
        }
    return {
        "package": "Audit Readiness Package",
        "coverage": "Coverage summary",
        "gaps": "Gap list",
        "maturity": "Maturity scores",
        "conflicts": "Confirmed conflicts",
        "code": "Code",
        "title": "Title",
        "covered": "Covered",
        "requirements": "Requirements",
        "supporting": "Supporting only",
        "doc": "Documentation",
        "impl": "Implementation",
        "scored": "Scored",
        "topic": "Topic",
        "difference": "Difference",
        "yes": "Yes",
        "no": "No",
    }


def _add_table(document: WordDocument, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for index, value in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = value
        _shade(cell, "17365D")
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            if row_index % 2 == 0:
                _shade(cells[index], "F4F7FA")
    _borders(table)


def _score(value: float | None) -> str:
    return "—" if value is None else str(value)


def build_readiness_docx(
    *,
    framework: Framework,
    coverage: list[CoverageRow],
    gaps: list[GapRow],
    maturity: MaturitySummaryOut | None,
    conflicts: list[PolicyConflict],
    language: str,
) -> bytes:
    labels = _labels(language)
    sections = _section_labels(language)
    name = framework.name_zh if language == "zh" else framework.name_en

    document = WordDocument()
    _configure(document)
    document.core_properties.title = f"{name} {sections['package']}"
    document.core_properties.author = "GRC Helper"
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(f"{name}\n{sections['package']}")
    document.add_paragraph(f"{labels['generated']}: {datetime.now(UTC).date().isoformat()}")

    document.add_heading(sections["coverage"], level=1)
    if coverage:
        _add_table(
            document,
            [sections["code"], sections["title"], sections["covered"], sections["requirements"]],
            [[row.code, row.title, str(row.covered), str(row.requirements)] for row in coverage],
        )
    else:
        document.add_paragraph(labels["none"])

    document.add_heading(sections["gaps"], level=1)
    if gaps:
        _add_table(
            document,
            [sections["code"], sections["title"], sections["supporting"]],
            [
                [row.code, row.title, sections["yes"] if row.has_supporting else sections["no"]]
                for row in gaps
            ],
        )
    else:
        document.add_paragraph(labels["none"])

    document.add_heading(sections["maturity"], level=1)
    if maturity is None:
        document.add_paragraph(labels["none"])
    else:
        overall = maturity.overall
        _add_table(
            document,
            [sections["scored"], sections["doc"], sections["impl"]],
            [[
                f"{overall.scored_items}/{overall.total_items} {labels['items']}",
                _score(overall.doc_average),
                _score(overall.impl_average),
            ]],
        )
        if maturity.groups:
            _add_table(
                document,
                [sections["code"], sections["title"], sections["doc"], sections["impl"]],
                [
                    [group.code, group.title, _score(group.doc_average), _score(group.impl_average)]
                    for group in maturity.groups
                ],
            )

    document.add_heading(sections["conflicts"], level=1)
    if conflicts:
        _add_table(
            document,
            [sections["topic"], sections["difference"]],
            [[row.topic, row.difference] for row in conflicts],
        )
    else:
        document.add_paragraph(labels["none"])

    output = BytesIO()
    document.save(output)
    return output.getvalue()
